#!/usr/bin/env python3
"""Combine all non-Jira CSV reports into a single multi-tab .xlsx workbook.

Runs each existing report (Projects, Users, Team Memberships, Teams, Custom Field
Usage, Portfolios) into a temporary directory — reusing their ``main()`` entry
points unchanged — then folds the resulting CSVs into one workbook with one tab
per report (the Portfolio report contributes two: summary + items). The Team→Project
Jira import CSV is intentionally excluded.

This is the only script that depends on a third-party library: ``openpyxl`` (the
Python stdlib has no native .xlsx writer). It is imported lazily so the module stays
importable — and the rest of the toolchain stays stdlib-only — without it installed.
Install with ``python3 -m pip install -r requirements.txt``.
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.lib.config import load_credentials
from src.lib.logging_utils import configure_logging, resolve_log_level
from scripts import (
    export_project_details_report,
    export_users,
    export_team_memberships,
    export_team_report,
    export_custom_field_usage_report,
    export_portfolio_report,
)


DEFAULT_DB_PATH = "data/asana.sqlite"
DEFAULT_OUTPUT_DIR = "exports"
DEFAULT_CONFIG_ROOT = "."
MAX_COLUMN_WIDTH = 60
SAFE_INT_DIGITS = 15  # Excel only keeps 15 significant digits; keep longer ids (gids) as text.
logger = logging.getLogger(__name__)


# Each report: (module, temp filename, accepts --workspace-gid, [(sheet title, produced filename), ...]).
REPORT_SPECS = [
    (export_project_details_report, "projects.csv", True, [("Projects", "projects.csv")]),
    (export_users, "users.csv", True, [("Users", "users.csv")]),
    (export_team_memberships, "team_memberships.csv", True, [("Team Memberships", "team_memberships.csv")]),
    (export_team_report, "teams.csv", False, [("Teams", "teams.csv")]),
    (export_custom_field_usage_report, "custom_field_usage.csv", True, [("Custom Field Usage", "custom_field_usage.csv")]),
    (
        export_portfolio_report,
        "portfolios.csv",
        True,
        [("Portfolios", "portfolios.csv"), ("Portfolio Items", "portfolios_items.csv")],
    ),
]


def parse_args(argv: Optional[Iterable[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", default=DEFAULT_DB_PATH, help="Path to the SQLite database (default: data/asana.sqlite).")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR, help="Directory to store the workbook (default: exports/).")
    parser.add_argument("--filename", help="Optional workbook filename. Defaults to all_reports_<timestamp>.xlsx.")
    parser.add_argument("--workspace-gid", help="Optional Asana workspace GID used to build URLs in the reports.")
    parser.add_argument("--log-level", default="INFO", help="Logging verbosity (default: INFO).")
    return parser.parse_args(argv)


def _coerce(value: str):
    """Turn CSV text into an int/float where safe, else keep the original string.

    Long all-digit strings (Asana GIDs exceed Excel's 15-digit precision) are kept
    as text so they are not silently corrupted. Blank cells become ``None``.
    """
    if value is None or value == "":
        return None
    if re.fullmatch(r"-?\d+", value):
        digits = value.lstrip("-")
        if len(digits) <= SAFE_INT_DIGITS and not (len(digits) > 1 and digits[0] == "0"):
            return int(value)
        return value
    if re.fullmatch(r"-?\d+\.\d+", value):
        try:
            return float(value)
        except ValueError:
            return value
    return value


def _safe_sheet_title(title: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", title).strip()
    return cleaned[:31] or "Sheet"


def _read_csv(path: Path) -> List[List[str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.reader(handle))


def _run_reports(database: str, tmp_dir: Path, workspace_gid: Optional[str], log_level: str) -> List[Tuple[str, List[List[str]]]]:
    """Run every report into ``tmp_dir`` and return [(sheet title, csv rows), ...]."""
    sheets: List[Tuple[str, List[List[str]]]] = []
    for module, filename, accepts_workspace, sheet_map in REPORT_SPECS:
        argv = ["--database", database, "--output-dir", str(tmp_dir), "--filename", filename, "--log-level", log_level]
        if accepts_workspace and workspace_gid:
            argv += ["--workspace-gid", workspace_gid]
        logger.info("Generating %s ...", ", ".join(title for title, _ in sheet_map))
        module.main(argv)
        for title, produced in sheet_map:
            sheets.append((title, _read_csv(tmp_dir / produced)))
    return sheets


def _write_workbook(sheets: Sequence[Tuple[str, List[List[str]]]], output_path: Path, database: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - exercised only without the dep
        raise SystemExit(
            "export_all_reports_xlsx requires openpyxl, which is not installed.\n"
            "Install it with: python3 -m pip install -r requirements.txt"
        ) from exc

    workbook = Workbook()
    workbook.remove(workbook.active)

    # Overview tab first: what's inside and when it was generated.
    overview = workbook.create_sheet(title="Overview")
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%SZ")
    overview.append(["Asana migration reports"])
    overview.append(["Generated (UTC)", generated_at])
    overview.append(["Source database", database])
    overview.append([])
    overview.append(["Tab", "Rows"])
    overview["A1"].font = Font(bold=True, size=14)
    overview["A5"].font = Font(bold=True)
    overview["B5"].font = Font(bold=True)

    for title, rows in sheets:
        sheet = workbook.create_sheet(title=_safe_sheet_title(title))
        data_row_count = max(len(rows) - 1, 0) if rows else 0
        overview.append([title, data_row_count])

        if not rows:
            sheet.append(["(no data)"])
            continue

        header, *body = rows
        sheet.append(list(header))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for record in body:
            sheet.append([_coerce(value) for value in record])
        sheet.freeze_panes = "A2"

        # Approximate column widths from header + a sample of rows.
        widths = [len(str(value)) for value in header]
        for record in body[:200]:
            for idx, value in enumerate(record):
                if idx < len(widths):
                    widths[idx] = max(widths[idx], len(str(value)))
        for idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(idx)].width = min(width + 2, MAX_COLUMN_WIDTH)

    for idx, width in enumerate((28, 22), start=1):
        overview.column_dimensions[get_column_letter(idx)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def export_all_reports_xlsx(
    database: str,
    output_dir: Path,
    filename: Optional[str] = None,
    workspace_gid: Optional[str] = None,
    log_level_name: str = "INFO",
) -> Path:
    output_dir = output_dir.resolve()
    if not filename:
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        filename = f"all_reports_{timestamp}.xlsx"
    output_path = output_dir / filename

    with tempfile.TemporaryDirectory() as tmp:
        sheets = _run_reports(database, Path(tmp), workspace_gid, log_level_name)
    _write_workbook(sheets, output_path, database)

    total_rows = sum(max(len(rows) - 1, 0) for _, rows in sheets)
    logger.info("Wrote %s tab(s) (%s data rows) to %s.", len(sheets), total_rows, output_path)
    return output_path


def main(argv: Optional[Iterable[str]] = None) -> None:
    args = parse_args(argv)
    try:
        log_level = resolve_log_level(args.log_level)
    except ValueError as exc:  # noqa: BLE001
        raise SystemExit(str(exc)) from exc
    configure_logging(log_level)

    workspace_gid = args.workspace_gid or load_credentials(DEFAULT_CONFIG_ROOT).workspace_gid
    output_path = export_all_reports_xlsx(
        database=args.database,
        output_dir=Path(args.output_dir),
        filename=args.filename,
        workspace_gid=workspace_gid,
        log_level_name=args.log_level,
    )
    logger.info("Combined reports workbook completed: %s", output_path)


if __name__ == "__main__":
    main()
