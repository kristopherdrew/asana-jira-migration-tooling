from __future__ import annotations

import importlib.util
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from unittest import TestCase, skipUnless

from src.lib.database import Database, synced_at_timestamp
from scripts.export_all_reports_xlsx import (
    _coerce,
    _safe_sheet_title,
    export_all_reports_xlsx,
)

OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None
SYNC = synced_at_timestamp(datetime(2026, 1, 2, tzinfo=timezone.utc))

EXPECTED_TABS = [
    "Overview",
    "Projects",
    "Users",
    "Team Memberships",
    "Teams",
    "Custom Field Usage",
    "Portfolios",
    "Portfolio Items",
]


class CoerceTest(TestCase):
    def test_blank_becomes_none(self) -> None:
        self.assertIsNone(_coerce(""))
        self.assertIsNone(_coerce(None))

    def test_small_integers_coerce(self) -> None:
        self.assertEqual(_coerce("0"), 0)
        self.assertEqual(_coerce("42"), 42)
        self.assertEqual(_coerce("-7"), -7)
        self.assertIsInstance(_coerce("42"), int)

    def test_floats_coerce(self) -> None:
        self.assertEqual(_coerce("0.42"), 0.42)
        self.assertIsInstance(_coerce("0.42"), float)

    def test_long_gid_stays_text(self) -> None:
        # 16-digit GID exceeds Excel's 15-digit precision; must remain a string.
        self.assertEqual(_coerce("1209238528105063"), "1209238528105063")
        self.assertIsInstance(_coerce("1209238528105063"), str)

    def test_leading_zero_stays_text(self) -> None:
        self.assertEqual(_coerce("007"), "007")

    def test_plain_text_passthrough(self) -> None:
        self.assertEqual(_coerce("on_track"), "on_track")
        self.assertEqual(_coerce("2026-01-02"), "2026-01-02")


class SafeSheetTitleTest(TestCase):
    def test_truncates_to_31_chars(self) -> None:
        self.assertEqual(len(_safe_sheet_title("x" * 40)), 31)

    def test_strips_illegal_characters(self) -> None:
        self.assertNotIn("/", _safe_sheet_title("a/b:c*d?"))

    def test_blank_falls_back(self) -> None:
        self.assertEqual(_safe_sheet_title("///"), "Sheet")


@skipUnless(OPENPYXL_AVAILABLE, "openpyxl not installed")
class ExportAllReportsXlsxTest(TestCase):
    def setUp(self) -> None:
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = str(Path(self.tmpdir.name) / "asana.sqlite")
        self.out_dir = Path(self.tmpdir.name) / "exports"
        db = Database(db_path=self.db_path)
        conn = db.connect()
        db.initialize(conn)
        db.upsert_user(conn, {"gid": "u_1", "name": "Olive", "email": "olive@x.com"}, SYNC)
        db.upsert_portfolio(
            conn,
            {"gid": "1209238528105063", "name": "Roadmap", "owner": {"gid": "u_1", "name": "Olive"}},
            SYNC,
        )
        conn.commit()
        conn.close()

    def tearDown(self) -> None:
        self.tmpdir.cleanup()

    def test_workbook_has_all_tabs_in_order(self) -> None:
        import openpyxl

        path = export_all_reports_xlsx(
            database=self.db_path, output_dir=self.out_dir, workspace_gid="ws_1", filename="all.xlsx"
        )
        self.assertTrue(path.exists())
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, EXPECTED_TABS)

    def test_overview_lists_row_counts(self) -> None:
        import openpyxl

        path = export_all_reports_xlsx(
            database=self.db_path, output_dir=self.out_dir, workspace_gid="ws_1"
        )
        wb = openpyxl.load_workbook(path)
        overview = {
            row[0]: row[1]
            for row in wb["Overview"].iter_rows(values_only=True)
            if row and row[0] in {"Portfolios", "Users"}
        }
        self.assertEqual(overview["Portfolios"], 1)
        self.assertEqual(overview["Users"], 1)

    def test_gid_kept_as_text_and_header_frozen(self) -> None:
        import openpyxl

        path = export_all_reports_xlsx(
            database=self.db_path, output_dir=self.out_dir, workspace_gid="ws_1"
        )
        wb = openpyxl.load_workbook(path)
        ws = wb["Portfolios"]
        header = [c.value for c in ws[1]]
        gid_col = header.index("portfolio_gid") + 1
        gid_cell = ws.cell(row=2, column=gid_col)
        self.assertEqual(gid_cell.value, "1209238528105063")
        self.assertIsInstance(gid_cell.value, str)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws["A1"].font.bold)
